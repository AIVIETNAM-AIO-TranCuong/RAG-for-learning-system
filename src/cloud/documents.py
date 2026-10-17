import hashlib
import re
import unicodedata
from pathlib import Path

from fastapi import HTTPException

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def safe_filename(value: str) -> str:
    # Handle both POSIX and Windows separators regardless of the server OS.
    value = value.replace("\\", "/").rsplit("/", 1)[-1]
    value = re.sub(r"[\x00-\x1f\x7f<>:\"|?*]", "_", value).strip(" .")
    if not value or len(value) > 240:
        raise HTTPException(422, "Tên file không hợp lệ hoặc quá dài.")
    if Path(value).suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise HTTPException(422, "Chỉ hỗ trợ PDF, DOCX, PPTX, TXT và Markdown.")
    return value


def clean_text(value: str) -> str:
    value = unicodedata.normalize("NFC", value)
    value = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", "", value)
    value = re.sub(r"[ \t]+", " ", value)
    return re.sub(r"\n{3,}", "\n\n", value).strip()


def logical_pages(parts: list[str], words=350):
    pages, current, count = [], [], 0
    for part in parts:
        if not part.strip():
            continue
        current.append(part)
        count += len(part.split())
        if count >= words:
            pages.append((len(pages) + 1, "\n\n".join(current)))
            current, count = [], 0
    if current:
        pages.append((len(pages) + 1, "\n\n".join(current)))
    return pages


def extract_chunks(path: Path, filename: str, document_id: str) -> list[dict]:
    extension = path.suffix.lower()
    try:
        if extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            pages = [(index, page.extract_text() or "") for index, page in enumerate(reader.pages, 1)]
        elif extension == ".docx":
            from docx import Document
            document = Document(path)
            parts = [p.text for p in document.paragraphs]
            parts += [" | ".join(cell.text for cell in row.cells)
                      for table in document.tables for row in table.rows]
            pages = logical_pages(parts)
        elif extension == ".pptx":
            from pptx import Presentation
            presentation = Presentation(path)
            pages = []
            for index, slide in enumerate(presentation.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text)
                    elif shape.has_table:
                        texts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                pages.append((index, "\n".join(texts)))
        else:
            data = path.read_bytes()
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1252", errors="replace")
            pages = logical_pages(text.split("\n\n"), words=400)
    except Exception as error:
        raise HTTPException(422, "Không đọc được file. Kiểm tra định dạng hoặc mật khẩu bảo vệ tài liệu.") from error

    chunks = []
    for page, value in pages:
        text = clean_text(value)
        start = 0
        while start < len(text):
            end = min(start + 1000, len(text))
            if end < len(text):
                boundary = text.rfind(" ", start + 700, end)
                if boundary > start:
                    end = boundary
            content = text[start:end].strip()
            if content:
                position = len(chunks)
                chunks.append(dict(document_id=document_id, filename=filename, source=filename,
                    page=page, chunk_id=f"{document_id}_{position}", position=position,
                    section=None, content=content))
            if end >= len(text):
                break
            start = end - 150
    if not chunks:
        raise HTTPException(422, "File không có nội dung chữ để trích xuất. PDF scan cần OCR trước khi upload.")
    return chunks


def document_digest(filename: str, content_hash: str) -> str:
    return hashlib.sha256(f"{filename}\0{content_hash}".encode()).hexdigest()[:32]
