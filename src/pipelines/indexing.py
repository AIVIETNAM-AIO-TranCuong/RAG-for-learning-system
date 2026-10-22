import hashlib
from pathlib import Path
from uuid import uuid4
from typing import List, Union, Optional, Any

from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from qdrant_client.http.models import PointStruct

from src.core.config import settings
from src.core.schemas import ChunkMetadata
from src.core.chunking import create_chunker
from src.db.store import get_store
from src.utils.tf_text_utils import preprocess_batch

# Danh sách các định dạng tài liệu được hỗ trợ
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def _make_document_id(file_path: Path) -> str:
    """Tạo ID duy nhất cho tài liệu dựa trên tên file và kích thước."""
    file_size = file_path.stat().st_size
    raw = f"{file_path.name}_{file_size}"
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def _load_pdf(path: Path) -> List[Document]:
    """Đọc tệp PDF theo từng trang."""
    try:
        loader = PyPDFLoader(str(path))
        return loader.load()
    except Exception as e:
        raise ValueError(f"Lỗi khi đọc file PDF '{path.name}': {str(e)}")


def _load_docx(path: Path) -> List[Document]:
    """Đọc tệp Word (.docx) và gom đoạn văn/bảng thành các trang tài liệu logic."""
    try:
        import docx
    except ImportError:
        raise ImportError("Thiếu thư viện 'python-docx'. Cài đặt bằng: pip install python-docx")

    try:
        doc = docx.Document(str(path))
    except Exception as e:
        raise ValueError(f"Lỗi khi mở file Word '{path.name}': {str(e)}")

    pages: List[Document] = []
    current_page_text: List[str] = []
    current_words = 0
    page_idx = 1

    # Đọc các đoạn văn (paragraphs)
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        current_page_text.append(text)
        current_words += len(text.split())
        # Cứ mỗi ~350-400 từ gom thành 1 trang logic để trích dẫn trang rõ ràng
        if current_words >= 350:
            pages.append(Document(
                page_content="\n\n".join(current_page_text),
                metadata={"page": page_idx, "source": str(path.resolve())}
            ))
            current_page_text = []
            current_words = 0
            page_idx += 1

    # Đọc các bảng biểu (tables)
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_rows.append(row_text)
        if table_rows:
            current_page_text.append("\n".join(table_rows))
            current_words += sum(len(r.split()) for r in table_rows)
            if current_words >= 350:
                pages.append(Document(
                    page_content="\n\n".join(current_page_text),
                    metadata={"page": page_idx, "source": str(path.resolve())}
                ))
                current_page_text = []
                current_words = 0
                page_idx += 1

    if current_page_text:
        pages.append(Document(
            page_content="\n\n".join(current_page_text),
            metadata={"page": page_idx, "source": str(path.resolve())}
        ))

    return pages


def _load_pptx(path: Path) -> List[Document]:
    """Đọc tệp PowerPoint (.pptx), mỗi slide được tính là 1 trang (page = slide number)."""
    try:
        import pptx
    except ImportError:
        raise ImportError("Thiếu thư viện 'python-pptx'. Cài đặt bằng: pip install python-pptx")

    try:
        prs = pptx.Presentation(str(path))
    except Exception as e:
        raise ValueError(f"Lỗi khi mở file PowerPoint '{path.name}': {str(e)}")

    pages: List[Document] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = paragraph.text.strip()
                    if t:
                        slide_texts.append(t)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_t = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_t:
                        slide_texts.append(row_t)
        if slide_texts:
            pages.append(Document(
                page_content="\n".join(slide_texts),
                metadata={"page": slide_idx, "source": str(path.resolve())}
            ))

    return pages


def _load_text(path: Path) -> List[Document]:
    """Đọc tệp văn bản thuần (.txt) hoặc Markdown (.md)."""
    encodings = ("utf-8", "utf-8-sig", "latin-1", "cp1252")
    content = ""
    for enc in encodings:
        try:
            with open(path, "r", encoding=enc) as f:
                content = f.read()
            break
        except UnicodeDecodeError:
            continue
    else:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()

    if not content.strip():
        return []

    # Gom theo đoạn thành các trang logic ~400 từ
    paragraphs = content.split("\n\n")
    pages: List[Document] = []
    current_page_text: List[str] = []
    current_words = 0
    page_idx = 1

    for p in paragraphs:
        p_clean = p.strip()
        if not p_clean:
            continue
        current_page_text.append(p_clean)
        current_words += len(p_clean.split())
        if current_words >= 400:
            pages.append(Document(
                page_content="\n\n".join(current_page_text),
                metadata={"page": page_idx, "source": str(path.resolve())}
            ))
            current_page_text = []
            current_words = 0
            page_idx += 1

    if current_page_text:
        pages.append(Document(
            page_content="\n\n".join(current_page_text),
            metadata={"page": page_idx, "source": str(path.resolve())}
        ))

    return pages


def load_file_to_documents(path: Path) -> List[Document]:
    """
    Dispatcher đọc nội dung tệp theo định dạng (PDF, DOCX, PPTX, TXT, MD).
    """
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path)
    elif ext == ".docx":
        return _load_docx(path)
    elif ext == ".pptx":
        return _load_pptx(path)
    elif ext in {".txt", ".md"}:
        return _load_text(path)
    else:
        raise ValueError(
            f"Định dạng tệp '{ext}' không được hỗ trợ. "
            f"Các định dạng hỗ trợ: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )


def index_document(
    file_path: Union[str, Path],
    force_reindex: bool = False,
    chunking_method: str = "recursive",
    chunk_size: Optional[int] = None,
    chunk_overlap: Optional[int] = None,
    semantic_threshold_type: str = "interquartile",
    semantic_threshold_amount: Optional[float] = None,
    store: Optional[Any] = None,
) -> str:
    """
    Đọc tài liệu đa định dạng (PDF, DOCX, PPTX, TXT, MD), tiền xử lý, chunking (recursive hoặc semantic),
    embedding và lưu vào Qdrant.
    
    Args:
        file_path: Đường dẫn đến tệp tài liệu
        force_reindex: Nếu True, xóa dữ liệu cũ và index lại nếu tài liệu đã tồn tại
        chunking_method: 'recursive' hoặc 'semantic'
        chunk_size: Kích thước chunk (nếu dùng recursive)
        chunk_overlap: Độ chồng lấp (nếu dùng recursive)
        semantic_threshold_type: 'interquartile', 'standard_deviation', hoặc 'percentile' (nếu dùng semantic)
        semantic_threshold_amount: Tham số tùy chỉnh ngưỡng semantic
        store: Custom QdrantStore instance (optional)
        
    Returns:
        document_id của tài liệu đã được index
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy tệp: {path}")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Định dạng tệp '{ext}' không được hỗ trợ. "
            f"Các định dạng hỗ trợ: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    # Thêm hậu tố method vào doc_id nếu là cấu hình tùy chỉnh để có thể index song song
    c_size = chunk_size or settings.chunk_size
    c_overlap = chunk_overlap or settings.chunk_overlap
    raw_doc_id = _make_document_id(path)
    if chunking_method == "semantic":
        document_id = f"{raw_doc_id}_sem_{semantic_threshold_type}"
    elif chunk_size is not None or chunk_overlap is not None:
        document_id = f"{raw_doc_id}_rec_{c_size}_{c_overlap}"
    else:
        document_id = raw_doc_id

    active_store = store or get_store()

    # Kiểm tra xem tài liệu đã từng được index chưa
    existing_records = active_store.scroll_by_document(document_id)
    if existing_records:
        if not force_reindex:
            print(f"Tài liệu {path.name} (ID: {document_id}) đã tồn tại trong {active_store.collection_name}. Bỏ qua index.")
            return document_id
        else:
            print(f"Tài liệu {path.name} (ID: {document_id}) đã tồn tại. Đang xóa dữ liệu cũ và index lại...")
            active_store.delete_document(document_id)

    # 1. LOAD TÀI LIỆU THEO ĐỊNH DẠNG
    print(f"Đang đọc tài liệu: {path.name} ({ext})...")
    pages = load_file_to_documents(path)

    if not pages:
        raise ValueError(f"Tệp {path.name} không có nội dung chữ để trích xuất.")

    # Lấy mảng văn bản thô
    raw_texts = [page.page_content for page in pages]

    # 2. PREPROCESS TEXT
    print("Tiền xử lý văn bản...")
    cleaned_texts = preprocess_batch(raw_texts)

    for i, page in enumerate(pages):
        page.page_content = cleaned_texts[i]

    # 3. CHUNKING TEXT
    if chunking_method == "semantic":
        print(f"Chunking ngữ nghĩa (Semantic Chunking, threshold={semantic_threshold_type})...")
        chunker = create_chunker(
            method="semantic",
            semantic_threshold_type=semantic_threshold_type,
            semantic_threshold_amount=semantic_threshold_amount,
            embeddings=active_store.embeddings
        )
    else:
        print(f"Chunking đệ quy (Recursive Chunking, chunk_size={c_size}, overlap={c_overlap})...")
        chunker = create_chunker(
            method="recursive",
            chunk_size=c_size,
            chunk_overlap=c_overlap
        )
    chunks = chunker.split_documents(pages)

    if not chunks:
        print("Cảnh báo: Không tạo được chunk nào từ tệp này.")
        return document_id

    # 4. EMBEDDING & CHUẨN BỊ VECTORS
    print(f"Đang tạo embedding và chuẩn bị lưu {len(chunks)} chunks vào Qdrant ({active_store.collection_name})...")
    points = []
    chunk_texts = [chunk.page_content for chunk in chunks]
    embeddings_list = active_store.embeddings.embed_documents(chunk_texts)

    for i, (chunk, embedding) in enumerate(zip(chunks, embeddings_list)):
        chunk_id = f"{document_id}_{i}"
        
        # Schema metadata chuẩn
        metadata = ChunkMetadata(
            document_id=document_id,
            filename=path.name,
            source=str(path.resolve()),
            page=chunk.metadata.get("page", 1),
            chunk_id=chunk_id,
            section=None
        )

        payload = metadata.model_dump()
        payload["content"] = chunk.page_content

        point = PointStruct(
            id=str(uuid4()),
            vector=embedding,
            payload=payload
        )
        points.append(point)

    # 5. UPSERT VÀO QDRANT
    print(f"Đang lưu {len(points)} chunks vào Qdrant...")
    active_store.upsert_chunks(points)
    print(f"Đã lập chỉ mục thành công tài liệu '{path.name}' với {len(chunks)} chunks vào {active_store.collection_name}.")

    return document_id


# Alias tương thích ngược với code cũ
index_pdf = index_document
