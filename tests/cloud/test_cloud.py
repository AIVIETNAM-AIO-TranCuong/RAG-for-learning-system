import asyncio
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from qdrant_client import AsyncQdrantClient, QdrantClient, models

from scripts.migrate_cloud import export_local
from src.cloud.ai import GeminiAI
from src.cloud.api import create_app
from src.cloud.config import CloudSettings, MAX_UPLOAD_BYTES
from src.cloud.documents import extract_chunks, safe_filename
from src.cloud.service import CloudService, bounded_context
from src.cloud.store import CloudStore, point_id


def settings(**overrides):
    return CloudSettings(_env_file=None, gemini_api_key="test-only-key", qdrant_api_key="test-only-key",
                         qdrant_url="https://example.qdrant.io", **overrides)


def chunks(count=3, document_id="doc-a"):
    return [dict(document_id=document_id, chunk_id=f"{document_id}_{i}", content=f"Nội dung chương {i} về RAG.",
                 filename=f"{document_id}.txt", source=f"{document_id}.txt", page=i + 1, position=i, section=None)
            for i in range(count)]


class FakeAI:
    def __init__(self):
        self.ranking = "[0, 1, 2]"
        self.prompts = []
        self.embedded = []
        self.active = self.peak = 0

    async def embed(self, text, **kwargs):
        self.embedded.append((text, kwargs))
        return [1.0] + [0.0] * 767

    async def embed_documents(self, values):
        return [await self.embed(value["content"], title=value["filename"]) for value in values]

    async def generate(self, prompt, json_output=False):
        self.prompts.append(prompt)
        if prompt.startswith("Xếp hạng"):
            if isinstance(self.ranking, Exception):
                raise self.ranking
            return self.ranking
        if "trac nghiem (MCQ)" in prompt:
            return json.dumps([dict(question=f"Câu hỏi {i}", options=["A. 1", "B. 2", "C. 3", "D. 4"], answer="A", explanation="Lý do") for i in range(20)])
        if "flashcards" in prompt:
            return json.dumps([dict(front=f"Thẻ {i}", back="Trả lời") for i in range(30)])
        if "tom tat" in prompt:
            self.active += 1
            self.peak = max(self.peak, self.active)
            try:
                await asyncio.sleep(0.001)
                return "Nội dung tóm tắt."
            finally:
                self.active -= 1
        return "Trả lời dựa trên tài liệu [S1]."


@pytest.fixture
async def service():
    store = CloudStore(settings())
    await store.client.close()
    store.client = AsyncQdrantClient(location=":memory:")
    service = CloudService(store, FakeAI())
    yield service
    await store.close()


def test_import_has_no_local_models():
    assert not {"torch", "tensorflow", "streamlit", "sentence_transformers", "src.db.store"} & sys_modules()


def test_cloud_settings_ignore_local_model_environment(monkeypatch):
    monkeypatch.setenv("EMBEDDING_MODEL", "intfloat/multilingual-e5-large")
    monkeypatch.setenv("GEMINI_MODEL", "local-original")
    cloud = settings()
    assert cloud.embedding_model == "gemini-embedding-2"
    assert cloud.model == "gemini-3.5-flash-lite"


def sys_modules():
    import sys
    return set(sys.modules)


def test_missing_configuration_is_503(monkeypatch):
    for key in ("GEMINI_API_KEY", "QDRANT_URL", "QDRANT_API_KEY"):
        monkeypatch.delenv(key, raising=False)
    config = CloudSettings(_env_file=None)
    with TestClient(create_app(config)) as client:
        response = client.get("/api/health")
        assert response.status_code == 503
        assert response.json()["status"] == "unavailable"
        assert client.get("/api/documents").status_code == 503


def test_built_ui_is_served_at_root_without_exposing_source(tmp_path):
    (tmp_path / "index.html").write_text('<h1>Learning studio</h1>', encoding="utf-8")
    (tmp_path / "app.mjs").write_text('export const ready = true;', encoding="utf-8")
    with TestClient(create_app(settings(), ui_directory=tmp_path)) as client:
        assert client.get("/").status_code == 200
        assert "Learning studio" in client.get("/").text
        assert client.get("/app.mjs").status_code == 200
        assert client.get("/.env.cloud").status_code == 404
        assert client.get("/src/cloud/api.py").status_code == 404


async def test_source_filter_citations_and_contracts(service):
    await service.index(chunks())
    await service.index(chunks(2, "doc-b"))
    app = create_app(settings(), service)
    # TestClient is safe here with an in-memory client; no network connection changes event loops.
    with TestClient(app) as client:
        assert client.get("/api/health").status_code == 200
        documents = client.get("/api/documents").json()
        assert len(documents) == 2
        result = client.post("/api/ask", json=dict(question="RAG là gì?", document_id="doc-a"))
        assert result.status_code == 200
        answer = result.json()
        assert all(c["metadata"]["document_id"] == "doc-a" for c in answer["chunks"])
        assert answer["citations"][0]["marker"] == "S1"
        assert client.post("/api/ask", json=dict(question="RAG là gì?", document_id=None)).status_code == 200
        assert client.post("/api/summarize", json={"document_id": "doc-a"}).json()["method"] == "single"
        quiz = client.post("/api/quiz", json=dict(document_id="doc-a", topic="RAG", n_questions=5))
        assert quiz.status_code == 200 and len(quiz.json()["items"]) == 5
        cards = client.post("/api/flashcards", json=dict(document_id="doc-a", n_cards=8))
        assert cards.status_code == 200 and len(cards.json()["cards"]) == 8
        assert client.delete("/api/documents/doc-a").status_code in (404, 405)
        for path, body in (("ask", {"question": "Test", "document_id": "missing"}),
                           ("summarize", {"document_id": "missing"}), ("quiz", {"document_id": "missing"}),
                           ("flashcards", {"document_id": "missing"})):
            assert client.post(f"/api/{path}", json=body).status_code == 404


@pytest.mark.parametrize("ranking", ["invalid JSON", "[0, 0, 1]", "[10, 1, 2]", "[true, 1, 2]", RuntimeError("provider unavailable")])
async def test_reranking_falls_back_to_retrieval(service, ranking):
    await service.index(chunks())
    service.ai.ranking = ranking
    answer = await service.ask("RAG?", "doc-a")
    assert len(answer.chunks) == 3 and answer.citations[0].marker == "S1"


async def test_pending_chunks_are_invisible_and_retry_does_not_duplicate(service, monkeypatch):
    original = service.store.client.set_payload
    monkeypatch.setattr(service.store.client, "set_payload", AsyncMock(side_effect=RuntimeError("commit failed")))
    with pytest.raises(RuntimeError):
        await service.index(chunks())
    assert await service.store.chunks("doc-a") == []
    assert await service.store.search([1.0] + [0.0] * 767, "doc-a") == []
    monkeypatch.setattr(service.store.client, "set_payload", original)
    await service.index(chunks())
    await service.index(chunks())
    assert len(await service.store.chunks("doc-a")) == 3
    assert (await service.store.client.count(service.store.settings.collection)).count == 3


async def test_embedding_failure_writes_no_chunks(service):
    service.ai.embed_documents = AsyncMock(side_effect=RuntimeError("embedding unavailable"))
    with pytest.raises(RuntimeError):
        await service.index(chunks())
    assert await service.store.documents() == []


async def test_map_reduce_covers_every_chunk_and_uses_at_most_two_calls(service):
    await service.index(chunks(31))
    result = await service.summarize("doc-a")
    assert result.method == "map_reduce" and result.total_chunks == 31
    assert service.ai.peak == 2
    map_prompts = [p for p in service.ai.prompts if "batch " in p]
    assert len(map_prompts) == 4
    for chunk in chunks(31):
        assert sum(chunk["content"] in p for p in map_prompts) == 1


def test_context_covers_beginning_and_end_within_budget():
    values = [dict(chunk, content=str(i) + "x" * 1000) for i, chunk in enumerate(chunks(300))]
    selected = bounded_context(values)
    assert sum(len(c["content"]) for c in selected) <= 48000
    assert selected[0]["chunk_id"] == "doc-a_0"
    assert selected[-1]["chunk_id"] == "doc-a_299"


async def test_bad_learning_json_is_502(service):
    await service.index(chunks())
    service.ai.generate = AsyncMock(return_value='[{"question":"incomplete"}]')
    with pytest.raises(HTTPException) as failure:
        await service.learning("quiz", "doc-a", None, 1)
    assert failure.value.status_code == 502


@pytest.mark.parametrize("status,expected", [(429, 429), (403, 502), (500, 502)])
async def test_provider_errors_do_not_expose_credentials(service, status, expected):
    class ProviderError(Exception):
        code = status
    service.store.documents = AsyncMock(side_effect=ProviderError("secret-key-and-payload"))
    with TestClient(create_app(settings(), service)) as client:
        response = client.get("/api/documents")
        assert response.status_code == expected
        assert "secret-key-and-payload" not in response.text


async def test_operation_timeout_is_504(service, monkeypatch):
    import src.cloud.api as api
    monkeypatch.setattr(api, "OPERATION_SECONDS", 0.01)
    service.store.documents = AsyncMock(side_effect=lambda: None)
    async def slow():
        await asyncio.sleep(1)
    service.store.documents = slow
    with TestClient(create_app(settings(), service)) as client:
        assert client.get("/api/documents").status_code == 504


async def test_upload_bounds_empty_duplicate_and_no_overwrite(service):
    with TestClient(create_app(settings(), service)) as client:
        assert client.post("/api/upload", files={"file": ("a.txt", b"x" * (MAX_UPLOAD_BYTES + 1))}).status_code == 413
        assert client.post("/api/upload", files={"file": ("a.txt", b"")}).status_code == 422
        assert client.post("/api/upload", files={"file": ("a.pdf", b"invalid PDF")}).status_code == 422
        assert client.post("/api/upload", files={"file": ("a.exe", b"test")}).status_code == 422
        assert client.post("/api/upload", files={"file": ("a.txt", b"test")}, data={"force_reindex": "true"}).status_code == 422
        for _ in range(2):
            result = client.post("/api/upload", files={"file": ("../../notes.txt", "Nội dung RAG".encode())})
            assert result.status_code == 200
            assert result.json()["filename"] == "notes.txt"
        assert len(await service.store.documents()) == 1
        assert len(await service.store.chunks()) == 1
        # Exact byte boundary is accepted before parsing/indexing.
        result = client.post("/api/upload", files={"file": ("boundary.txt", b"x" * MAX_UPLOAD_BYTES)})
        assert result.status_code == 200


@pytest.mark.parametrize("extension", [".pdf", ".docx", ".pptx", ".txt", ".md"])
def test_five_file_formats_and_citation_pages(tmp_path, extension):
    path = tmp_path / ("notes" + extension)
    if extension == ".pdf":
        # A minimal text PDF with correct xref offsets, independent from any PDF writer dependency.
        objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
        stream = b"BT /F1 12 Tf 20 200 Td (RAG document text) Tj ET"
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        data, offsets = b"%PDF-1.4\n", [0]
        for i, obj in enumerate(objects, 1):
            offsets.append(len(data)); data += str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
        xref = len(data)
        data += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets[1:])
        data += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF"
        path.write_bytes(data)
    elif extension == ".docx":
        from docx import Document
        doc = Document(); doc.add_paragraph("RAG document text"); doc.save(path)
    elif extension == ".pptx":
        from pptx import Presentation
        ppt = Presentation(); slide = ppt.slides.add_slide(ppt.slide_layouts[1]); slide.shapes.title.text = "RAG document text"; ppt.save(path)
    else:
        path.write_text("RAG document text", encoding="utf-8")
    result = extract_chunks(path, path.name, "test-document")
    assert result[0]["page"] == 1 and "RAG document text" in result[0]["content"]


def test_snapshot_is_read_only_and_preserves_payload(tmp_path):
    source = tmp_path / "original"
    client = QdrantClient(path=str(source))
    client.create_collection("documents", vectors_config=models.VectorParams(size=2, distance=models.Distance.COSINE))
    expected = dict(chunks(1)[0], source="C:/private/local/path/notes.txt")
    client.upsert("documents", points=[models.PointStruct(id=point_id(expected), vector=[1, 0], payload=expected)])
    source_db = source / "collection/documents/storage.sqlite"
    before = hashlib.sha256(source_db.read_bytes()).digest()
    exported = export_local(source)  # Original Qdrant lock is still held.
    assert exported[0]["source"] == "doc-a.txt"
    assert exported[0]["content"] == expected["content"]
    assert hashlib.sha256(source_db.read_bytes()).digest() == before
    client.close()


async def test_sdk_embeds_each_chunk_separately():
    ai = object.__new__(GeminiAI)
    ai.settings = settings()
    sdk = AsyncMock(return_value=SimpleNamespace(embeddings=[SimpleNamespace(values=[0.1] * 768)]))
    ai.client = SimpleNamespace(aio=SimpleNamespace(models=SimpleNamespace(embed_content=sdk)))
    result = await ai.embed_documents(chunks())
    assert len(result) == 3 and sdk.await_count == 3
    assert all(isinstance(call.kwargs["contents"], str) for call in sdk.await_args_list)
    assert "title: doc-a.txt | text:" in sdk.await_args_list[0].kwargs["contents"]
