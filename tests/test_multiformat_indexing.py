import sys
from pathlib import Path
REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

"""
Test suite cho tinh nang ho tro da dinh dang tai lieu:
- Microsoft Word (.docx)
- PowerPoint (.pptx)
- Plain Text (.txt)
- Markdown (.md)
- PDF (.pdf)
"""

import sys
import tempfile
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import docx
import pptx
from src.pipelines.indexing import (
    index_document,
    load_file_to_documents,
    SUPPORTED_EXTENSIONS,
)
from src.db.store import get_store
from src.pipelines.rag import retrieve


def test_docx_loading_and_indexing():
    print("--- [1] Testing Word (.docx) Loading & Indexing ---")
    with tempfile.TemporaryDirectory() as tmp_dir:
        doc_path = Path(tmp_dir) / "test_machine_learning.docx"
        doc = docx.Document()
        doc.add_heading("Machine Learning Overview", level=1)
        doc.add_paragraph("Machine learning is a subset of artificial intelligence focused on building applications that learn from data and improve their accuracy over time without being programmed to do so.")
        doc.add_paragraph("Supervised learning is where you have input variables (x) and an output variable (Y) and you use an algorithm to learn the mapping function from the input to the output.")
        doc.save(str(doc_path))

        pages = load_file_to_documents(doc_path)
        assert len(pages) > 0, "Docx pages should not be empty"
        assert "Machine learning" in pages[0].page_content

        doc_id = index_document(doc_path, force_reindex=True)
        assert doc_id is not None and len(doc_id) == 16
        print(f"✅ Docx indexed successfully with ID: {doc_id}")

        # Test retrieval
        chunks = retrieve("What is supervised learning?", top_k=3, document_id=doc_id)
        assert len(chunks) > 0
        assert any("Supervised learning" in c.content for c in chunks)
        print("✅ Docx retrieval verified!")


def test_pptx_loading_and_indexing():
    print("\n--- [2] Testing PowerPoint (.pptx) Loading & Indexing ---")
    with tempfile.TemporaryDirectory() as tmp_dir:
        ppt_path = Path(tmp_dir) / "test_deep_learning_slides.pptx"
        prs = pptx.Presentation()
        
        # Slide 1
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Deep Learning Architecture"
        slide1.placeholders[1].text = "Convolutional Neural Networks (CNN) are widely used for image recognition and computer vision tasks."
        
        # Slide 2
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Transformer Architecture"
        slide2.placeholders[1].text = "Transformers rely entirely on self-attention mechanisms to compute representations of its input and output without using sequence-aligned RNNs or convolution."
        
        prs.save(str(ppt_path))

        pages = load_file_to_documents(ppt_path)
        assert len(pages) == 2, f"Expected 2 slides/pages, got {len(pages)}"
        assert pages[0].metadata["page"] == 1
        assert pages[1].metadata["page"] == 2
        assert "Transformer" in pages[1].page_content

        doc_id = index_document(ppt_path, force_reindex=True)
        assert doc_id is not None
        print(f"✅ PPTX indexed successfully with ID: {doc_id}")

        # Test retrieval
        chunks = retrieve("self-attention mechanism in transformers", top_k=3, document_id=doc_id)
        assert len(chunks) > 0
        assert any("Transformers" in c.content for c in chunks)
        print("✅ PPTX retrieval verified!")


def test_txt_and_md_loading_and_indexing():
    print("\n--- [3] Testing Text (.txt) & Markdown (.md) Loading & Indexing ---")
    with tempfile.TemporaryDirectory() as tmp_dir:
        # TXT
        txt_path = Path(tmp_dir) / "test_notes.txt"
        txt_path.write_text("Reinforcement Learning is an aspect of Machine learning where an agent learns to behave in an environment by performing actions and seeing the results.", encoding="utf-8")
        
        pages_txt = load_file_to_documents(txt_path)
        assert len(pages_txt) > 0
        doc_id_txt = index_document(txt_path, force_reindex=True)
        print(f"✅ TXT indexed successfully with ID: {doc_id_txt}")

        # MD
        md_path = Path(tmp_dir) / "test_guide.md"
        md_path.write_text("# Vector Databases\nVector databases index and store vector embeddings for fast retrieval and similarity search, such as Qdrant.", encoding="utf-8")
        
        pages_md = load_file_to_documents(md_path)
        assert len(pages_md) > 0
        doc_id_md = index_document(md_path, force_reindex=True)
        print(f"✅ MD indexed successfully with ID: {doc_id_md}")

        # Test retrieval
        chunks = retrieve("Qdrant similarity search", top_k=2, document_id=doc_id_md)
        assert len(chunks) > 0
        assert any("Qdrant" in c.content for c in chunks)
        print("✅ MD retrieval verified!")


def test_unsupported_format():
    print("\n--- [4] Testing Unsupported File Format Rejection ---")
    with tempfile.TemporaryDirectory() as tmp_dir:
        bad_path = Path(tmp_dir) / "test_script.exe"
        bad_path.write_text("binary dummy", encoding="utf-8")
        try:
            index_document(bad_path)
            assert False, "Should have raised ValueError"
        except ValueError as e:
            assert "không được hỗ trợ" in str(e)
            print(f"✅ Unsupported format properly rejected: {e}")


if __name__ == "__main__":
    print("🚀 Bắt đầu kiểm thử toàn diện Multi-format Indexing Pipeline...")
    test_docx_loading_and_indexing()
    test_pptx_loading_and_indexing()
    test_txt_and_md_loading_and_indexing()
    test_unsupported_format()
    print("\n🎉 Tất cả bài kiểm tra Multi-format Indexing ĐÃ VƯỢT QUA 100%!")
